import calendar
from datetime import date

from django.contrib.auth import get_user_model, login
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .forms import SignUpForm
from .models import AttendanceRecord

User = get_user_model()
STATUS_LABELS = dict(AttendanceRecord.Status.choices)
VALID_STATUSES = set(STATUS_LABELS.keys())
BOLD = Font(bold=True)


def register(request):
    if request.user.is_authenticated:
        return redirect("tracker")
    if request.method == "POST":
        form = SignUpForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect("tracker")
    else:
        form = SignUpForm()
    return render(request, "tracker/register.html", {"form": form})


@login_required
def tracker_view(request):
    """
    Regular users: editable calendar with their own records embedded.
    Staff users: reports panel + a READ-ONLY calendar of a selected user.
    """
    today = date.today()
    context = {
        "current_year": today.year,
        "current_month": today.month,  # 1-indexed for the template
    }
    if request.user.is_staff:
        staff_users = User.objects.filter(is_staff=False).order_by("username")
        selected_user = None
        raw = request.GET.get("view_user", "")
        if raw:
            if not raw.isdigit():
                return HttpResponseBadRequest("Invalid user.")
            selected_user = get_object_or_404(staff_users, pk=int(raw))
        context["staff_users"] = staff_users
        context["months"] = [(i, calendar.month_name[i]) for i in range(1, 13)]
        context["selected_user"] = selected_user
        context["status_labels"] = STATUS_LABELS
        context["records"] = (
            {
                r.date.isoformat(): r.status
                for r in AttendanceRecord.objects.filter(user=selected_user)
            }
            if selected_user
            else {}
        )
    else:
        context["records"] = {
            r.date.isoformat(): r.status
            for r in AttendanceRecord.objects.filter(user=request.user)
        }
    return render(request, "tracker/tracker.html", context)


@login_required
@require_POST
def set_status(request):
    """Create, update, or clear a single day's status for the signed-in (non-staff) user only."""
    if request.user.is_staff:
        raise PermissionDenied  # staff can only download reports, not fill the calendar
    day = request.POST.get("date", "")
    status = request.POST.get("status", "")

    try:
        parsed_date = date.fromisoformat(day)
    except ValueError:
        return HttpResponseBadRequest("Invalid date.")

    if status == "":
        AttendanceRecord.objects.filter(user=request.user, date=parsed_date).delete()
        return JsonResponse({"ok": True, "date": day, "status": None})

    if status not in VALID_STATUSES:
        return HttpResponseBadRequest("Invalid status.")

    AttendanceRecord.objects.update_or_create(
        user=request.user, date=parsed_date, defaults={"status": status}
    )
    return JsonResponse({"ok": True, "date": day, "status": status})


def _resolve_target(request):
    """
    Return a User, or None meaning "all users" (staff only).
    Non-staff users can only ever get their own data.
    Assumes user_id was already validated as "", "all", or digits.
    """
    raw = request.GET.get("user_id", "")
    if raw in ("", str(request.user.pk)):
        return request.user
    if not request.user.is_staff:
        raise PermissionDenied
    if raw == "all":
        return None
    return get_object_or_404(User, pk=int(raw))


@login_required
def download_report(request):
    """Export attendance as .xlsx (month or full year). Staff may pick any user or all users."""
    scope = request.GET.get("scope")

    try:
        year = int(request.GET.get("year", ""))
        if not 1900 <= year <= 2100:
            raise ValueError
    except ValueError:
        return HttpResponseBadRequest("A valid year is required.")

    raw_user = request.GET.get("user_id", "")
    if raw_user not in ("", "all") and not raw_user.isdigit():
        return HttpResponseBadRequest("Invalid user.")
    target = _resolve_target(request)  # 403 for non-staff, 404 if user missing

    qs = AttendanceRecord.objects.filter(date__year=year)
    if target is not None:
        qs = qs.filter(user=target)
        who = target.username
    else:
        qs = qs.select_related("user")
        who = "all_users"

    if scope == "month":
        try:
            month = int(request.GET.get("month", ""))
            if not 1 <= month <= 12:
                raise ValueError
        except ValueError:
            return HttpResponseBadRequest("A valid month is required.")
        qs = qs.filter(date__month=month)
        filename = f"attendance_{who}_{year}_{month:02d}.xlsx"
        sheet_title = "Month Report"
    elif scope == "year":
        filename = f"attendance_{who}_{year}_full_year.xlsx"
        sheet_title = "Year Report"
    else:
        return HttpResponseBadRequest("scope must be 'month' or 'year'.")

    include_user_col = target is None
    header = ["Date", "Day", "Month", "Status"]
    widths = [12, 12, 12, 18]
    if include_user_col:
        header.insert(0, "User")
        widths.insert(0, 18)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD

    counts = {}
    order = ("user__username", "date") if include_user_col else ("date",)
    for record in qs.order_by(*order):
        label = STATUS_LABELS.get(record.status, record.status)
        row = [
            record.date.isoformat(),
            record.date.strftime("%A"),
            record.date.strftime("%B"),
            label,
        ]
        if include_user_col:
            row.insert(0, record.user.username)
        ws.append(row)
        counts[label] = counts.get(label, 0) + 1

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    if scope == "year":
        summary = wb.create_sheet("Summary")
        summary.append(["Status", "Days"])
        for cell in summary[1]:
            cell.font = BOLD
        for label, count in counts.items():
            summary.append([label, count])
        summary.column_dimensions["A"].width = 20
        summary.column_dimensions["B"].width = 8

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response